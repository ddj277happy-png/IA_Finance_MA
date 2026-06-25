"""Parse IPAI (housing price index) Excel → CSV."""
import openpyxl
import re
from pathlib import Path
import pandas as pd

SRC = Path(r"D:/ProjectforMM/IA_Finance_MA01/data/_raw_pdf/BKAM_Séries_IPAI_T1_2026.xlsx")
OUT_DIR = Path(r"D:/ProjectforMM/IA_Finance_MA01/data/_intermediate")
OUT_DIR.mkdir(parents=True, exist_ok=True)

FR_TO_MONTH = {"T1": 1, "T2": 4, "T3": 7, "T4": 10}


def parse_trimestre(s):
    """Parse 'T12006' or 'T1 2018' or 'T1-2025' → Timestamp."""
    if not isinstance(s, str):
        return None
    m = re.match(r"T\s*([1-4])\s*[- ]?\s*(\d{4})", s)
    if not m:
        return None
    q, y = int(m.group(1)), int(m.group(2))
    return pd.Timestamp(year=y, month=FR_TO_MONTH[f"T{q}"], day=1)


wb = openpyxl.load_workbook(SRC, data_only=True)
ws = wb["Global"]

# First row is title; second row is header
header = [c.value for c in ws[2]]
print("Header:", header)

rows = []
for row in ws.iter_rows(min_row=3, values_only=True):
    if not row[0]:
        continue
    date = parse_trimestre(row[0])
    if date is None:
        continue
    rec = {"date": date}
    for col_idx, col_name in enumerate(header[1:], start=1):
        val = row[col_idx] if col_idx < len(row) else None
        if val is not None and isinstance(val, (int, float)):
            rec[col_name.strip()] = float(val)
    rows.append(rec)

df = pd.DataFrame(rows).set_index("date").sort_index()
print(f"\nShape: {df.shape}")
print(f"Range: {df.index.min()} -> {df.index.max()}")
print(df.tail(8))

out = OUT_DIR / "ipai_quarterly.csv"
df.to_csv(out)
print(f"\nSaved: {out}")