"""Parse BAM IPAI (housing price index) Excel file."""
import openpyxl
from pathlib import Path
import pandas as pd

fp = Path(r"D:/ProjectforMM/IA_Finance_MA01/data/_raw_pdf/BKAM_Séries_IPAI_T1_2026.xlsx")
print(f"File: {fp.name} ({fp.stat().st_size} bytes)")

wb = openpyxl.load_workbook(fp, data_only=True)
print(f"Sheets: {wb.sheetnames}")
for sn in wb.sheetnames:
    ws = wb[sn]
    print(f"\n=== Sheet '{sn}' ({ws.max_row} rows x {ws.max_column} cols) ===")
    for i, row in enumerate(ws.iter_rows(values_only=True)):
        if i < 10 or i > ws.max_row - 5:
            print(f"  R{i+1}: {row}")
        elif i == 10:
            print("  ...")